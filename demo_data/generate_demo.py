#!/usr/bin/env python3

from fpdf import FPDF
from docx import Document
import os

# Fictional data
jd_content = """
Senior Python Developer

We are looking for a skilled Senior Python Developer with experience in web development and data analysis.

Required Skills:
- Python (3+ years)
- Django or Flask
- SQL and NoSQL databases
- Machine Learning with scikit-learn
- API development

3+ years of experience preferred.
"""

cv1_content = """
John Doe
Experience:
- 5 years in Python development
- Worked with Django, Flask
- Proficient in SQL, MySQL
- Experience in data science and ML

Skills: Python, Django, SQL, Machine Learning, REST APIs
"""

cv2_content = """
Jane Smith
Experience:
- 2 years in web development
- Knowledge of Python, JavaScript
- Worked with React and Node.js
- Basic database knowledge

Skills: Python, JavaScript, HTML, CSS, React
"""

cv3_content = """
Bob Johnson
Experience:
- 4 years in data analysis
- Proficient in Python, SQL
- Experience with pandas, numpy
- ETL processes

Skills: Python, SQL, pandas, numpy, ETL
"""

files = [
    ("JD.pdf", jd_content),
    ("JD.docx", jd_content),
    ("John_Doe_CV.pdf", cv1_content),
    ("John_Doe_CV.docx", cv1_content),
    ("Jane_Smith_CV.pdf", cv2_content),
    ("Jane_Smith_CV.docx", cv2_content),
    ("Bob_Johnson_CV.pdf", cv3_content),
    ("Bob_Johnson_CV.docx", cv3_content),
]


def create_pdf(filename, content):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    # Add content line by line
    for line in content.split('\n'):
        pdf.cell(200, 10, txt=line, ln=True)
    pdf.output(filename)


def create_docx(filename, content):
    doc = Document()
    for line in content.split('\n'):
        if line.strip():
            doc.add_paragraph(line)
    doc.save(filename)


if __name__ == "__main__":
    os.chdir(os.path.dirname(__file__))
    for filename, content in files:
        if filename.endswith('.pdf'):
            create_pdf(filename, content.strip())
        elif filename.endswith('.docx'):
            create_docx(filename, content.strip())
    print("Demo files generated!")
